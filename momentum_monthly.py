"""Monthly momentum ETF schedule, cloud version of the old monthly task's steps 1-6.

Runs every day as part of the morning job. Does nothing unless (a) today is on or after the
first London Stock Exchange trading day of the month and (b) the workbook does not already hold
this month's schedule. That way a missed day is caught up the next time the job runs.

Prices: month-end closes from Yahoo daily data. Checked on 5 October 2026 against the
stockanalysis.com prices already in prices.csv for June to September 2026: all 19 columns
agreed to within 0.01%. Unemployment: FRED UNRATE through the FRED API.

The rules and the calculation are untouched: this only refreshes prices.csv and
unemployment.csv and then runs momentum_schedule.py exactly as before.

Prints one line starting STATUS: not-due / produced / failed, then detail.
Exit code 0 unless it failed.
"""
import csv, json, math, os, shutil, ssl, subprocess, sys, urllib.request
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo

HERE = os.path.dirname(os.path.abspath(__file__))
MDIR = os.environ.get("MOMENTUM_DIR", os.path.join(HERE, "momentum"))
CTX = ssl.create_default_context()
UA = {"User-Agent": "Mozilla/5.0"}
YAHOO = {"CEMQ": "CEMQ.DE", "CEMS": "CEMS.DE", "SPY": "SPY"}   # everything else is TICKER.L
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August",
          "September", "October", "November", "December"]


def get_json(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=40, context=CTX) as r:
        return json.load(r)


def easter(y):
    a = y % 19; b = y // 100; c = y % 100; d = b // 4; e = b % 4; f = (b + 8) // 25
    g = (b - f + 1) // 3; h = (19 * a + b - d - g + 15) % 30; i = c // 4; k = c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7; m = (a + 11 * h + 22 * l) // 451
    mo = (h + l - 7 * m + 114) // 31; da = ((h + l - 7 * m + 114) % 31) + 1
    return date(y, mo, da)


def bank_holidays(y):
    def sub(d):                      # weekend holiday moves to the next free weekday
        while d.weekday() >= 5:
            d += timedelta(days=1)
        return d
    def first_mon(m):
        d = date(y, m, 1)
        return d + timedelta(days=(7 - d.weekday()) % 7)
    def last_mon(m):
        d = date(y, m + 1, 1) - timedelta(days=1) if m < 12 else date(y, 12, 31)
        return d - timedelta(days=d.weekday())
    e = easter(y)
    xmas = date(y, 12, 25); box = date(y, 12, 26)
    if xmas.weekday() == 5:   xmas_o, box_o = date(y, 12, 27), date(y, 12, 28)   # Sat: Mon and Tue
    elif xmas.weekday() == 6: xmas_o, box_o = date(y, 12, 27), box               # Sun: Tue for Christmas, Mon is Boxing Day
    elif box.weekday() == 5:  xmas_o, box_o = xmas, date(y, 12, 28)              # Fri Christmas: Boxing Day moves to Mon
    else:                     xmas_o, box_o = xmas, box
    return {sub(date(y, 1, 1)), e - timedelta(days=2), e + timedelta(days=1), first_mon(5),
            last_mon(5), last_mon(8), xmas_o, box_o}


def trading_day(d):
    return d.weekday() < 5 and d not in bank_holidays(d.year)


def first_trading_day(y, m):
    d = date(y, m, 1)
    while not trading_day(d):
        d += timedelta(days=1)
    return d


def fmt7(v):
    """Seven significant figures, as the old task wrote them (JavaScript toPrecision(7))."""
    if v == 0:
        return "0.000000"
    digits = max(0, 7 - (int(math.floor(math.log10(abs(v)))) + 1))
    return f"{v:.{digits}f}"


def workbook_month(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    title = str(wb["This month"].cell(1, 1).value or "")
    for i, name in enumerate(MONTHS):
        if f"for {name} " in title:
            return f"{title.split()[-1]}-{i + 1:02d}", title
    return None, title


def main():
    today = datetime.now(ZoneInfo("Europe/London")).date()
    cur = f"{today.year:04d}-{today.month:02d}"
    ftd = first_trading_day(today.year, today.month)
    wbp = os.path.join(MDIR, "ETF schedule.xlsx")
    have, title = workbook_month(wbp)
    if today < ftd:
        print(f"STATUS: not-due\nFirst trading day of {MONTHS[today.month-1]} is {ftd}; workbook holds {title!r}")
        return 0
    if have == cur and not os.environ.get("MOMENTUM_FORCE"):   # MOMENTUM_FORCE: testing only
        print(f"STATUS: not-due\nWorkbook already holds {title!r}")
        return 0

    prev = (date(today.year, today.month, 1) - timedelta(days=1))
    M = f"{prev.year:04d}-{prev.month:02d}"
    with open(os.path.join(MDIR, "prices.csv"), newline="") as f:
        rows = list(csv.reader(f))
    head = rows[0]
    old = {r[0]: r for r in rows[1:] if r}

    # month-end closes from Yahoo daily data
    month_end, last_date = {}, {}
    for t in head[1:]:
        sym = YAHOO.get(t, t + ".L")
        j = get_json(f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}?range=2y&interval=1d")
        r = j["chart"]["result"][0]
        closes = r["indicators"]["quote"][0]["close"]
        tz = ZoneInfo(r.get("meta", {}).get("exchangeTimezoneName") or "UTC")
        per = {}
        for ts, c in zip(r["timestamp"], closes):
            if c is None:
                continue
            d = datetime.fromtimestamp(ts, tz).date()
            per[f"{d.year:04d}-{d.month:02d}"] = (c, d)
        month_end[t] = {k: v[0] for k, v in per.items()}
        last_date[t] = per.get(M, (None, None))[1]
    months = sorted(k for k in month_end[head[1]] if k < cur)[-14:]

    problems = []
    if not months or months[-1] != M:
        problems.append(f"latest complete month from Yahoo is {months[-1] if months else 'none'}, expected {M}")
    for t in head[1:]:
        if M not in month_end[t]:
            problems.append(f"{t}: no price for {M}")
        elif last_date[t] is None or last_date[t] < prev - timedelta(days=7):
            problems.append(f"{t}: last price in {M} is from {last_date[t]}, not the last week of the month")
    for mo in months:
        if mo in old and mo != M:
            for i, t in enumerate(head[1:], start=1):
                if old[mo][i] and mo in month_end[t]:
                    o = float(old[mo][i]); n = month_end[t][mo]
                    if o and abs(n / o - 1) > 0.01:
                        problems.append(f"{t} {mo}: new price {n:.4f} differs from recorded {o:.4f} by more than 1%")
    if problems:
        print("STATUS: failed\nNo schedule produced. " + "; ".join(problems[:12]))
        return 1

    for mo in months:
        old[mo] = [mo] + [fmt7(month_end[t][mo]) if mo in month_end[t] else "" for t in head[1:]]

    # US unemployment, FRED API
    key = os.environ.get("FRED_API_KEY")
    u = get_json("https://api.stlouisfed.org/fred/series/observations?series_id=UNRATE"
                 f"&api_key={key}&file_type=json&observation_start={today.year - 3}-01-01")
    with open(os.path.join(MDIR, "unemployment.csv"), newline="") as f:
        urows = list(csv.reader(f))
    uh = urows[0]; uold = {r[0]: r for r in urows[1:] if r}
    for o in u.get("observations", [])[-26:]:
        mo = o["date"][:7]
        uold[mo] = [mo, "" if o["value"] in ("", ".") else o["value"]]

    # keep copies, then write and run the unchanged calculation
    for fn in ("prices.csv", "unemployment.csv", "ETF schedule.xlsx"):
        shutil.copy2(os.path.join(MDIR, fn), os.path.join(MDIR, fn + ".before"))
    with open(os.path.join(MDIR, "prices.csv"), "w", newline="") as f:
        w = csv.writer(f, lineterminator="\n"); w.writerow(head)
        for mo in sorted(old):
            w.writerow(old[mo])
    with open(os.path.join(MDIR, "unemployment.csv"), "w", newline="") as f:
        w = csv.writer(f, lineterminator="\n"); w.writerow(uh)
        for mo in sorted(uold):
            w.writerow(uold[mo])
    res = subprocess.run([sys.executable, "momentum_schedule.py", "prices.csv", "unemployment.csv",
                          "ETF schedule.xlsx"], cwd=MDIR, capture_output=True, text=True)
    if res.returncode != 0 or f"Schedule for {cur}" not in res.stdout:
        for fn in ("prices.csv", "unemployment.csv", "ETF schedule.xlsx"):
            shutil.copy2(os.path.join(MDIR, fn + ".before"), os.path.join(MDIR, fn))
        print("STATUS: failed\nThe calculation did not produce this month's schedule; inputs restored.\n"
              + (res.stdout + res.stderr)[-1500:])
        return 1
    print("STATUS: produced\n" + res.stdout)
    return 0


if __name__ == "__main__":
    sys.exit(main())
